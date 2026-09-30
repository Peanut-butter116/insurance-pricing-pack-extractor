from pptx import Presentation


def segments(path):
    output = []
    for i, slide in enumerate(Presentation(path).slides, 1):
        for j, shape in enumerate(slide.shapes, 1):
            if shape.has_table:
                for k, row in enumerate(shape.table.rows, 1):
                    output.append({'location': f'Slide {i}, table {j}, row {k}',
                                   'text': ' | '.join(c.text for c in row.cells)})
            elif shape.has_text_frame and shape.text.strip():
                output.append({'location': f'Slide {i}, shape {j}', 'text': shape.text.strip()})
    return output
