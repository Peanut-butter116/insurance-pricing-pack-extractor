from datetime import datetime, date
from openpyxl import load_workbook


def display(cell):
    value = cell.value
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, (int, float)):
        if '%' in cell.number_format:
            return f'{value * 100:g}%'
        if '£' in cell.number_format or 'GBP' in cell.number_format:
            return f'£{value:,.2f}'
    return str(value)


def segments(path):
    book = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    output = []
    try:
        for sheet in book:
            for row in sheet:
                cells = [c for c in row if c.value is not None]
                for c in row:
                    if formulas[sheet.title][c.coordinate].data_type == 'f' and c.value is None:
                        raise ValueError(f'Uncached formula at {sheet.title}!{c.coordinate}; recalculate in Excel')
                if cells:
                    output.append({'location': f'Sheet {sheet.title}, cells ' + ', '.join(c.coordinate for c in cells),
                                   'text': ' | '.join(display(c) for c in cells)})
    finally:
        book.close()
        formulas.close()
    return output
