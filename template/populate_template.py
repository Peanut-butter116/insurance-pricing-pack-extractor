from datetime import date
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import ROOT, FIELDS, LABELS
from validation.schema_check import check_values


def create_template():
    path = ROOT / 'template/pricing_template.xlsx'
    if path.exists():
        return path
    book = Workbook()
    sheet = book.active
    sheet.title = 'Pricing review'
    sheet.merge_cells('A1:D1')
    sheet['A1'] = 'Pricing pack | Analyst review'
    sheet['A1'].font = Font(size=20, color='FFFFFF', bold=True)
    sheet['A1'].fill = PatternFill('solid', fgColor='16324F')
    sheet.row_dimensions[1].height = 40
    sheet.merge_cells('A2:D2')
    sheet['A2'] = 'Extracted information only. Analyst approval required. No pricing decisions.'
    sheet.append(['Field', 'Value', 'Source', 'Validation status'])
    for i, label in enumerate(LABELS, 4):
        sheet.cell(i, 1, label)
        sheet.row_dimensions[i].height = 52
        for c in sheet[i]:
            c.alignment = Alignment(wrap_text=True, vertical='center')
            c.fill = PatternFill('solid', fgColor='EDF2F7' if i % 2 == 0 else 'FFFFFF')
    for col, width in [('A', 26), ('B', 23), ('C', 65), ('D', 42)]:
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = 'C4'
    sheet.sheet_view.showGridLines = False
    sheet.print_area = 'A1:D11'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    book.save(path)
    return path


def populate(extraction, destination, source=''):
    values = check_values({field: extraction[field]['value'] for field in FIELDS})
    book = load_workbook(create_template())
    sheet = book.active
    for i, field in enumerate(FIELDS, 4):
        evidence, value = extraction[field], values[field]
        c = sheet.cell(i, 2)
        if value is None:
            c.value = None
        elif field.startswith('policy_'):
            c.value, c.number_format = date.fromisoformat(value), 'dd/mm/yyyy'
        elif field == 'rate_change_pct':
            c.value, c.number_format = value / 100, '0.0%'
        else:
            c.value = value
            c.number_format = '#,##0' if field == 'claim_count' else '"£"#,##0.00'
        # Explicit string types prevent uploaded evidence becoming Excel formulas.
        sheet.cell(i, 3, '\n'.join(x for x in [source, evidence.get('source_location'), evidence.get('source_text')] if x)).data_type = 's'
        sheet.cell(i, 4, evidence.get('validation_status', 'Review required') if value is not None else evidence.get('validation_status', 'Not found')).data_type = 's'
    book.save(destination)
    return destination
