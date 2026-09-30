"""Portable Python benchmark generator. Labels are never embedded as hidden metadata."""
from datetime import date
from config import ROOT, FIELDS, LABELS
from data.generate_ground_truth import ensure


def rendered(record, field):
    value = record['values'][field]
    if field in record['partial_dates']:
        return record['partial_dates'][field]
    if value is None:
        return None
    i = (record['variant'] + FIELDS.index(field)) % 5
    if field.startswith('policy_'):
        return date.fromisoformat(value).strftime(['%d/%m/%Y', '%d %B %Y', '%Y-%m-%d'][i % 3])
    if field == 'claim_count':
        return str(value)
    if field == 'rate_change_pct':
        if value == 7.5 and record['variant'] % 9 == 0:
            return 'seven and a half percent'
        return f'{value:+g}%' if i % 2 else f'{value:g} per cent'
    return [f'£{value:,}', f'£{value/1e6:g}m', f'GBP {value/1e6:g} million',
            f'{value/1e6:g} million GBP', f'£{value/1000:g}k'][i]


def entries(record):
    layout = record['layout']
    labels = list(LABELS)
    if layout == 'C':
        labels = ['Inception', 'Expiry', 'GWP', 'Reported claims', 'Total incurred', 'Rate movement', 'Annual aggregate', 'Each and every loss']
    if layout == 'D':
        labels = ['Cover begins', 'Cover ends', 'Gross written premium', 'Claims count', 'Claims incurred', 'Rate movement', 'Programme limit', 'Per occurrence']
    sentences = ['Cover begins on {}.', 'Cover ends on {}.', 'The gross written premium was {}.',
                 'Reported claims: {}.', 'Total incurred losses were {}.', 'Rate change: {}.',
                 'The aggregate programme limit is {}.', 'The per occurrence limit is {}.']
    order = [6, 7, 2, 5, 3, 4, 0, 1] if layout == 'D' else list(range(8))
    output = []
    for i in order:
        field = FIELDS[i]
        text = rendered(record, field)
        if text is not None:
            output.append((field, labels[i], text, sentences[i].format(text)))
    return output


def pdf(record, path):
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors
    styles = getSampleStyleSheet()
    story = [Paragraph('Commercial portfolio | Pricing summary', styles['Title']),
             Paragraph('SYNTHETIC RESEARCH DOCUMENT - UK commercial insurance', styles['Normal']), Spacer(1, 24)]
    rows = entries(record)
    layout = record['layout']
    def table(items, reverse=False):
        cells = [[Paragraph(x[2] if reverse else x[1], styles['Normal']), Paragraph(x[1] if reverse else x[2], styles['Normal'])] for x in items]
        t = Table(cells, colWidths=[225, 225])
        t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#edf2f7')),
                               ('BOTTOMPADDING', (0, 0), (-1, -1), 13), ('TOPPADDING', (0, 0), (-1, -1), 13),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP')]))
        return t
    if layout == 'A':
        story.append(table(rows))
    elif layout == 'B':
        for x in rows:
            story += [Paragraph(x[3], styles['BodyText']), Spacer(1, 17)]
    elif layout == 'C':
        for heading, fields in [('Policy information', FIELDS[:2]), ('Premium and movement', (FIELDS[2], FIELDS[5])),
                                ('Claims experience', FIELDS[3:5]), ('Limit profile', FIELDS[6:])]:
            story.append(Paragraph(heading, styles['Heading2']))
            selected = [x for x in rows if x[0] in fields]
            if selected:
                story.append(table(selected))
    else:
        story += [Paragraph('Renewal dashboard', styles['Heading1']), table(rows[:4], reverse=True), Spacer(1, 20)]
        for x in rows[4:]:
            story += [Paragraph(f'{x[1]}: {x[2]}', styles['BodyText']), Spacer(1, 13)]
    story += [Spacer(1, 20), Paragraph('Prepared for analyst review. No pricing recommendation is provided.', styles['Normal'])]
    SimpleDocTemplate(str(path), title='Synthetic pricing summary', author='Synthetic research generator').build(story)


def excel(record, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    book = Workbook()
    book.remove(book.active)
    layout = record['layout']
    rows = entries(record)
    groups = [('Portfolio', rows)] if layout != 'C' else [('Policy', [x for x in rows if x[0] in FIELDS[:3]]), ('Experience', [x for x in rows if x[0] not in FIELDS[:3]])]
    for title, items in groups:
        sheet = book.create_sheet(title)
        sheet.merge_cells('A1:E1')
        sheet['A1'] = 'SYNTHETIC | Commercial portfolio pricing summary'
        sheet['A1'].font = Font(size=16, bold=True, color='FFFFFF')
        sheet['A1'].fill = PatternFill('solid', fgColor='16324F')
        sheet.row_dimensions[1].height = 32
        for col in 'ABCDE':
            sheet.column_dimensions[col].width = 27
        for idx, (field, label, text, sentence) in enumerate(items, 4):
            row = idx * 2 if layout == 'D' else idx
            if layout == 'B':
                sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
                sheet.cell(row, 1, sentence)
            else:
                lc, vc = (4, 2) if layout == 'D' else (1, 3)
                sheet.cell(row, lc, label)
                c = sheet.cell(row, vc, text)
                value = record['values'][field]
                if value is not None and record['variant'] % 2 == 0:
                    if field.startswith('policy_'):
                        c.value, c.number_format = date.fromisoformat(value), 'dd/mm/yyyy'
                    elif field == 'rate_change_pct':
                        c.value, c.number_format = value / 100, '0.0%'
                    else:
                        c.value = value
                        c.number_format = '#,##0' if field == 'claim_count' else '"£"#,##0'
            sheet.row_dimensions[row].height = 30
            for cell in sheet[row]:
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                if row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='EDF2F7')
        sheet.freeze_panes = 'A4'
        sheet.sheet_view.showGridLines = False
        sheet.print_options.horizontalCentered = True
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
    book.save(path)


def powerpoint(record, path):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.33), Inches(7.5)
    rows = entries(record)
    def box(slide, text, x, y, w, h, size=24):
        frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        frame.word_wrap = True
        frame.text = text
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = RGBColor.from_string('16324F')
    chunks = [rows[i:i+4] for i in range(0, len(rows), 4)]
    for index, items in enumerate(chunks):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box(slide, 'Commercial portfolio' if index == 0 else 'Portfolio detail', .6, .3, 12, .6, 32)
        box(slide, 'SYNTHETIC RESEARCH DOCUMENT', .6, 6.8, 12, .4, 14)
        layout = record['layout']
        if layout == 'A':
            table = slide.shapes.add_table(len(items), 2, Inches(.6), Inches(1.5), Inches(12), Inches(4.5)).table
            for i, x in enumerate(items):
                table.cell(i, 0).text, table.cell(i, 1).text = x[1], x[2]
                for cell in table.rows[i].cells:
                    for p in cell.text_frame.paragraphs:
                        p.font.size = Pt(24)
        elif layout == 'D':
            for i, x in enumerate(items):
                box(slide, x[2] + '\n' + x[1], .7 + (i % 2) * 6.3, 1.5 + (i // 2) * 2.3, 5.7, 1.8, 25)
        else:
            for i, x in enumerate(items):
                text = x[3] if layout == 'B' else f'{x[1]}: {x[2]}'
                box(slide, text, .7, 1.5 + i * 1.2, 11.9, 1.05, 23)
    deck.save(path)


def generate():
    records = ensure()
    for record in records:
        path = ROOT / record['path']
        if not path.exists():
            path.parent.mkdir(parents=True, exist_ok=True)
            {'pdf': pdf, 'excel': excel, 'powerpoint': powerpoint}[record['format']](record, path)
    return records


if __name__ == '__main__':
    print(f'{len(generate())} documents ready')
